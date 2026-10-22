# -*- coding: utf-8 -*-
"""Project, stage-run and approval routes.

Thin transport layer: validation here, state-machine rules in
app/services/flow.py (single implementation, §5).
"""
import base64
import io
from pathlib import Path
from typing import Annotated, Literal

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from pydantic import BaseModel, Field
from sqlalchemy import or_, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.api.deps import get_current_user
from app.config import get_settings
from app.db import get_db
from app.models import (
    ApprovalStatus,
    ApprovalTask,
    ArtifactFormat,
    ArtifactVersion,
    Project,
    ProjectMember,
    ProjectRole,
    ProjectStatus,
    StageKey,
    StageRun,
    StageRunStatus,
    User,
)
from app.services import flow
from app.services.flow import FlowError
from app.runtime_client import RuntimeCallError, get_runtime_client

router = APIRouter(tags=["projects"])


# ---------- schemas ----------

class ProjectCreateIn(BaseModel):
    name: str = Field(min_length=2, max_length=200)


class ProjectOut(BaseModel):
    id: int
    name: str
    status: ProjectStatus
    current_stage: StageKey
    owner_user_id: int
    created_at: str

    model_config = {"from_attributes": True}

    @classmethod
    def from_project(cls, project: Project) -> "ProjectOut":
        return cls(
            id=project.id,
            name=project.name,
            status=project.status,
            current_stage=project.current_stage,
            owner_user_id=project.owner_user_id,
            created_at=project.created_at.isoformat(),
        )


class StageRunOut(BaseModel):
    id: int
    project_id: int
    stage: StageKey
    status: StageRunStatus
    round: int
    session_id: str | None
    created_at: str

    model_config = {"from_attributes": True}

    @classmethod
    def from_run(cls, run: StageRun) -> "StageRunOut":
        return cls(
            id=run.id,
            project_id=run.project_id,
            stage=run.stage,
            status=run.status,
            round=run.round,
            session_id=run.session_id,
            created_at=run.created_at.isoformat(),
        )


class StartStageOut(StageRunOut):
    """StageRunOut plus why the runtime session may not be ready yet."""

    warning: str | None = None


class ChatIn(BaseModel):
    content: str = Field(min_length=1, max_length=20_000)


class ConfirmIn(BaseModel):
    """HITL verdict for the run's asking tool calls.

    ``tool_call_ids`` empty/None = every currently-asking call.
    ``allow_all``（仅 confirmed=true 时生效）额外把会话切到
    ``accept_edits``：工作目录内的后续文件读写不再逐个询问。
    """

    confirmed: bool = True
    tool_call_ids: list[str] | None = None
    allow_all: bool = False


class SubmitLatestIn(BaseModel):
    title: str = Field(min_length=1, max_length=200)


class SubmitFileIn(BaseModel):
    """Pick a file out of the agent's session workspace to submit.

    ``path`` is workspace-root-relative; the endpoint rejects traversal
    (``..``, absolute paths, backslashes) and non-whitelisted suffixes
    — belt and braces on top of the runtime's own session-root
    confinement.
    """

    title: str = Field(min_length=1, max_length=200)
    path: str = Field(min_length=1, max_length=500)


class WorkspaceFileOut(BaseModel):
    path: str
    name: str
    size_bytes: int | None = None


class MessagesOut(BaseModel):
    messages: list[dict]
    is_running: bool
    has_more: bool = False
    event_cursor: str | None = None


class SubmitIn(BaseModel):
    title: str = Field(min_length=1, max_length=200)
    content: str = Field(min_length=1, max_length=200_000)


class ArtifactOut(BaseModel):
    id: int
    version_no: int
    title: str
    format: str
    storage_path: str


class DecideIn(BaseModel):
    action: str = Field(pattern=r"^(approve|return)$")
    comment: str | None = Field(default=None, max_length=4000)


class ApprovalOut(BaseModel):
    id: int
    stage_run_id: int
    project_id: int
    project_name: str
    stage: StageKey
    run_round: int
    run_status: StageRunStatus
    status: ApprovalStatus
    comment: str | None


class MemberOut(BaseModel):
    """A membership row joined with the user's name fields."""

    id: int
    user_id: int
    username: str
    display_name: str
    role: ProjectRole
    created_at: str


class MemberAddIn(BaseModel):
    """Owner-only: add a member by username. INITIATOR is reserved for
    the owner (auto-created with the project), so it's not choosable."""

    username: str = Field(min_length=2, max_length=64)
    role: Literal["manager", "member", "approver", "observer"]


# ---------- helpers ----------

def _flow_http(exc: FlowError | RuntimeCallError) -> HTTPException:
    return HTTPException(status_code=exc.status, detail=exc.detail)


async def _get_run_or_404(
    db: AsyncSession, run_id: int, user: User
) -> tuple[StageRun, Project]:
    run = await db.get(StageRun, run_id)
    if run is None:
        raise HTTPException(status_code=404, detail="stage run not found")
    project = await _get_project_or_404(db, run.project_id, user)
    return run, project


async def _get_project_or_404(
    db: AsyncSession, project_id: int, user: User
) -> Project:
    project = await db.get(Project, project_id)
    if project is None:
        raise HTTPException(status_code=404, detail="project not found")
    try:
        await flow.require_member(db, project, user)
    except FlowError as exc:
        raise _flow_http(exc) from exc
    return project


# ---------- projects ----------

@router.post("/projects", response_model=ProjectOut, status_code=201)
async def create_project(
    payload: ProjectCreateIn,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> ProjectOut:
    project = Project(
        name=payload.name,
        owner_user_id=user.id,
    )
    db.add(project)
    await db.flush()
    db.add(
        ProjectMember(
            project_id=project.id,
            user_id=user.id,
            role=ProjectRole.INITIATOR,
        )
    )
    await db.commit()
    await db.refresh(project)
    return ProjectOut.from_project(project)


@router.get("/projects", response_model=list[ProjectOut])
async def list_projects(
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> list[ProjectOut]:
    member_project_ids = select(ProjectMember.project_id).where(
        ProjectMember.user_id == user.id
    )
    projects = (
        await db.scalars(
            select(Project)
            .where(
                or_(
                    Project.owner_user_id == user.id,
                    Project.id.in_(member_project_ids),
                )
            )
            .order_by(Project.id.desc())
        )
    ).all()
    return [ProjectOut.from_project(p) for p in projects]


@router.get("/projects/{project_id}", response_model=ProjectOut)
async def get_project(
    project_id: int,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> ProjectOut:
    project = await _get_project_or_404(db, project_id, user)
    return ProjectOut.from_project(project)


# ---------- project members ----------
# Approval addressees resolve at submit time (flow.resolve_approver,
# Decision D8): a member added as APPROVER receives every task submitted
# afterwards; tasks already addressed keep their original approver.


@router.get(
    "/projects/{project_id}/members", response_model=list[MemberOut]
)
async def list_members(
    project_id: int,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> list[MemberOut]:
    await _get_project_or_404(db, project_id, user)
    rows = (
        await db.execute(
            select(ProjectMember, User)
            .join(User, ProjectMember.user_id == User.id)
            .where(ProjectMember.project_id == project_id)
            .order_by(ProjectMember.id)
        )
    ).all()
    return [
        MemberOut(
            id=member.id,
            user_id=u.id,
            username=u.username,
            display_name=u.display_name,
            role=member.role,
            created_at=member.created_at.isoformat(),
        )
        for member, u in rows
    ]


@router.post(
    "/projects/{project_id}/members", response_model=MemberOut, status_code=201
)
async def add_member(
    project_id: int,
    payload: MemberAddIn,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> MemberOut:
    project = await _get_project_or_404(db, project_id, user)
    if project.owner_user_id != user.id:
        raise HTTPException(
            status_code=403, detail="only the project owner can manage members"
        )
    target = await db.scalar(
        select(User).where(User.username == payload.username)
    )
    if target is None:
        raise HTTPException(
            status_code=404,
            detail=f"user '{payload.username}' not found — they need to "
            f"register first",
        )
    if target.id == project.owner_user_id:
        raise HTTPException(
            status_code=409, detail="owner is already the project initiator"
        )
    exists = await db.scalar(
        select(ProjectMember).where(
            ProjectMember.project_id == project_id,
            ProjectMember.user_id == target.id,
        )
    )
    if exists is not None:
        raise HTTPException(status_code=409, detail="already a member")
    member = ProjectMember(
        project_id=project_id,
        user_id=target.id,
        role=ProjectRole(payload.role),
    )
    db.add(member)
    await db.commit()
    await db.refresh(member)
    return MemberOut(
        id=member.id,
        user_id=target.id,
        username=target.username,
        display_name=target.display_name,
        role=member.role,
        created_at=member.created_at.isoformat(),
    )


@router.delete(
    "/projects/{project_id}/members/{user_id}", status_code=204
)
async def remove_member(
    project_id: int,
    user_id: int,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> None:
    project = await _get_project_or_404(db, project_id, user)
    if project.owner_user_id != user.id:
        raise HTTPException(
            status_code=403, detail="only the project owner can manage members"
        )
    if user_id == project.owner_user_id:
        raise HTTPException(
            status_code=409, detail="cannot remove the owner's initiator row"
        )
    member = await db.scalar(
        select(ProjectMember).where(
            ProjectMember.project_id == project_id,
            ProjectMember.user_id == user_id,
        )
    )
    if member is None:
        raise HTTPException(status_code=404, detail="member not found")
    await db.delete(member)
    await db.commit()


# ---------- stage runs ----------

@router.get(
    "/projects/{project_id}/stage-runs", response_model=list[StageRunOut]
)
async def list_stage_runs(
    project_id: int,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> list[StageRunOut]:
    await _get_project_or_404(db, project_id, user)
    runs = (
        await db.scalars(
            select(StageRun)
            .where(StageRun.project_id == project_id)
            .order_by(StageRun.id)
        )
    ).all()
    return [StageRunOut.from_run(r) for r in runs]


@router.post(
    "/projects/{project_id}/stages/{stage}/start",
    response_model=StartStageOut,
    status_code=201,
)
async def start_stage(
    project_id: int,
    stage: StageKey,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> StartStageOut:
    """Open the next run of *stage* — DB only, the agent is NOT fired.

    The runtime session and the kickoff context (upstream artifacts,
    return comment) wait for the run's first chat message, so between
    审批/开启 and 开始 the user has a window to add an opening note
    and materials (e.g. the 功能点明细 xlsx for the budget stage).
    """
    project = await _get_project_or_404(db, project_id, user)
    try:
        run = await flow.start_stage(db, project, stage, user)
    except FlowError as exc:
        raise _flow_http(exc) from exc
    return StartStageOut.from_run(run)


@router.post("/stage-runs/{run_id}/submit", response_model=ArtifactOut)
async def submit_run(
    run_id: int,
    payload: SubmitIn,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> ArtifactOut:
    run, project = await _get_run_or_404(db, run_id, user)
    try:
        artifact = await flow.submit_run(
            db, run, project, payload.title, payload.content
        )
    except FlowError as exc:
        raise _flow_http(exc) from exc
    return ArtifactOut(
        id=artifact.id,
        version_no=artifact.version_no,
        title=artifact.title,
        format=artifact.format.value,
        storage_path=artifact.storage_path,
    )


# ---------- agent conversation (relayed through sys:runtime) ----------

@router.get("/stage-runs/{run_id}/artifacts", response_model=list[ArtifactOut])
async def list_run_artifacts(
    run_id: int,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> list[ArtifactOut]:
    run, _project = await _get_run_or_404(db, run_id, user)
    artifacts = (
        await db.scalars(
            select(ArtifactVersion)
            .where(ArtifactVersion.stage_run_id == run.id)
            .order_by(ArtifactVersion.version_no)
        )
    ).all()
    return [
        ArtifactOut(
            id=a.id,
            version_no=a.version_no,
            title=a.title,
            format=a.format.value,
            storage_path=a.storage_path,
        )
        for a in artifacts
    ]


@router.get("/stage-runs/{run_id}/artifacts/{artifact_id}")
async def get_run_artifact(
    run_id: int,
    artifact_id: int,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> dict:
    """Artifact content for preview (path is server-generated, never
    client-supplied, so no traversal surface)."""
    run, _project = await _get_run_or_404(db, run_id, user)
    artifact = await db.get(ArtifactVersion, artifact_id)
    if artifact is None or artifact.stage_run_id != run.id:
        raise HTTPException(status_code=404, detail="artifact not found")
    content = (
        Path(get_settings().artifact_dir) / artifact.storage_path
    ).read_text(encoding="utf-8")
    return {
        "id": artifact.id,
        "version_no": artifact.version_no,
        "title": artifact.title,
        "format": artifact.format.value,
        "content": content,
    }



async def _relay_chat(
    db: AsyncSession,
    run: StageRun,
    project: Project,
    message: str,
    blocks: list[dict] | None = None,
) -> str:
    """Shared tail of the two chat relays: ensure the session, self-heal
    the model config, deliver the message. Returns the session id.

    The FIRST message into a run's session doubles as the kickoff:
    opening a stage does not fire the agent, so the stage context
    (all upstream approved artifacts, plus the return comment on a
    rework round) rides along here.

    The model PATCH is skipped when a run holds the session (409): a
    running reply already carries the config it was started with, and
    the runtime refuses mid-run changes anyway.

    A parked session (HITL asking / external result) only accepts its
    answer event — a plain message there crashes the runtime run with
    a generic setup error and an empty bubble, so it is blocked with a
    clear 409 instead.
    """
    client = get_runtime_client()
    try:
        session_id, agent_id = await flow.ensure_stage_session(db, run, project)
        status = await client.get_session_status(agent_id, session_id)
        if status in ("awaiting_permission", "awaiting_external_result"):
            raise HTTPException(
                status_code=409,
                detail="智能体正在等待工具调用的确认，请先在对话中"
                "批准或拒绝（或中止本轮）后再发新消息",
            )
        # Empty projection on an idle session = no kickoff has landed
        # yet — including the self-heal case where an earlier attempt
        # failed right after creating the session.
        data = await client.get_messages(agent_id, session_id)
        if not data["messages"]:
            message = (
                f"{await flow.build_kickoff(db, run, project)}"
                f"\n\n【用户开场补充】\n{message}"
            )
        model = await client.default_chat_model_config()
        if model is None:
            raise HTTPException(
                status_code=503,
                detail="平台尚未配置模型凭证（sys:runtime 名下没有 "
                "credential），请联系管理员或用引导脚本配置后再对话",
            )
        try:
            await client.update_session_model(agent_id, session_id, model)
        except RuntimeCallError as exc:
            if exc.status != 409:
                raise
        await client.send_message(agent_id, session_id, message, blocks)
    except FlowError as exc:
        raise _flow_http(exc) from exc
    except RuntimeCallError as exc:
        if exc.status == 409:
            raise HTTPException(
                status_code=409,
                detail="智能体正在处理上一条消息，请等它完成后再发",
            ) from exc
        raise _flow_http(exc) from exc
    return session_id


@router.post("/stage-runs/{run_id}/chat")
async def chat_with_stage_run(
    run_id: int,
    payload: ChatIn,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> dict:
    """Relay a user message into the run's runtime session.

    Creates the session lazily (covers runs opened while the runtime
    was down) and PATCHes the model config on every send, so a
    credential registered after the session self-heals in.
    """
    run, project = await _get_run_or_404(db, run_id, user)
    if run.status is not StageRunStatus.IN_PROGRESS:
        raise HTTPException(
            status_code=409,
            detail=f"run #{run.id} is {run.status.value}, "
            f"chat is only open while in_progress",
        )
    session_id = await _relay_chat(db, run, project, payload.content)
    return {"ok": True, "session_id": session_id}


# ---------- chat with file upload ----------

# Text-like files are inlined as TextBlocks so any model can read
# them; images and PDFs travel as base64 DataBlocks (the runtime's
# OpenAI-compatible formatters map them to image_url / file_data).
# Word documents and Excel workbooks are neither — the formatters
# would silently DROP a docx/xlsx DataBlock — but they are the most
# common requirement/costing inputs in a bank (the 功能点明细 for
# the budget stage arrives as xlsx from the user's own external
# model), so .docx/.xlsx are extracted to text here (same inline
# path as md/txt). Everything else (zip/legacy .doc/.xls...) is
# still rejected with a clear message instead of a silent drop.
_TEXT_INLINE_EXTS = {
    ".md", ".txt", ".csv", ".json", ".log", ".xml", ".yaml", ".yml",
    ".py", ".js", ".ts", ".tsx", ".java", ".sql", ".html", ".htm",
    ".css", ".sh", ".bat", ".ini", ".toml",
}
_IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".webp"}
_TEXT_INLINE_LIMIT = 512 * 1024  # decoded characters
_UPLOAD_LIMIT = 8 * 1024 * 1024  # raw bytes


def _docx_to_text(raw: bytes) -> str:
    """Extract a .docx body as readable markdown-ish text.

    Paragraphs keep their heading level (style names are English or
    localized "标题 N"), tables become pipe tables in document order.
    Images cannot be conveyed on the text path and are reported as
    skipped so the agent (and user) know content was left behind.
    """
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(raw))
    parts: list[str] = []
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            para = Paragraph(child, doc)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name or "").lower()
            if style.startswith("heading") or style.startswith("标题"):
                level = "".join(c for c in style if c.isdigit())
                parts.append("#" * min(int(level or 1), 6) + " " + text)
            else:
                parts.append(text)
        elif child.tag.endswith("}tbl"):
            table = Table(child, doc)
            rows = [
                [cell.text.strip().replace("\n", " ") for cell in row.cells]
                for row in table.rows
            ]
            rows = [r for r in rows if any(r)]
            if not rows:
                continue
            # pad ragged rows (merged cells) to one column count
            width = max(len(r) for r in rows)
            rows = [r + [""] * (width - len(r)) for r in rows]
            parts.append(
                "| " + " | ".join(rows[0]) + " |\n"
                + "| " + " | ".join(["---"] * width) + " |\n"
                + "\n".join("| " + " | ".join(r) + " |" for r in rows[1:])
            )
    if len(doc.inline_shapes) > 0:
        parts.append(f"（注：文档内 {len(doc.inline_shapes)} 张图片无法以文本提取，已跳过）")
    return "\n\n".join(parts)


_XLSX_ROW_CAP = 500  # rows per sheet; bigger sheets truncate with a note


def _xlsx_to_text(raw: bytes) -> str:
    """Extract an .xlsx workbook as markdown pipe tables, one per sheet.

    Row one of each sheet is treated as the header; trailing blank
    rows (exports pad with them) are trimmed; merged cells come out
    as empty; a monster sheet truncates instead of blowing up the
    inline text block.
    """
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        while rows and all(c is None or not str(c).strip() for c in rows[-1]):
            rows.pop()
        if not rows:
            continue
        total = len(rows)
        if total > _XLSX_ROW_CAP:
            rows = rows[:_XLSX_ROW_CAP]
        width = max(len(r) for r in rows)
        cells = [
            [
                ""
                if c is None
                else str(c).replace("|", "\\|").replace("\n", " ")
                for c in r
            ]
            + [""] * (width - len(r))
            for r in rows
        ]
        table = (
            "| " + " | ".join(cells[0]) + " |\n"
            + "| " + " | ".join(["---"] * width) + " |\n"
            + "\n".join("| " + " | ".join(r) + " |" for r in cells[1:])
        )
        if total > _XLSX_ROW_CAP:
            table += f"\n\n（注：该工作表共 {total} 行，仅提取前 {_XLSX_ROW_CAP} 行）"
        parts.append(f"## 工作表：{ws.title}\n\n{table}")
    return "\n\n".join(parts)


@router.post("/stage-runs/{run_id}/chat-upload")
async def chat_with_upload(
    run_id: int,
    file: Annotated[UploadFile, File()],
    text: Annotated[str | None, Form(max_length=20_000)] = None,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> dict:
    """Relay a user message with one attached file into the run's
    session — the requirement agent reading an existing需求文档, the
    testing agent reading source code, an approver pasting a
    screenshot all go through here."""
    run, project = await _get_run_or_404(db, run_id, user)
    if run.status is not StageRunStatus.IN_PROGRESS:
        raise HTTPException(
            status_code=409,
            detail=f"run #{run.id} is {run.status.value}, "
            f"chat is only open while in_progress",
        )

    name = file.filename or "upload.bin"
    raw = await file.read()
    if not raw:
        raise HTTPException(status_code=415, detail="文件为空")
    if len(raw) > _UPLOAD_LIMIT:
        raise HTTPException(
            status_code=413,
            detail=f"文件超过 {_UPLOAD_LIMIT // (1024 * 1024)}MB 上限",
        )
    suffix = Path(name).suffix.lower()
    media = file.content_type or ""

    blocks: list[dict]
    if suffix in _TEXT_INLINE_EXTS or media.startswith("text/"):
        try:
            decoded = raw.decode("utf-8")
        except UnicodeDecodeError:
            raise HTTPException(
                status_code=415,
                detail="该文件不是 UTF-8 文本，请另存为 md/txt 后重试",
            ) from None
        if len(decoded) > _TEXT_INLINE_LIMIT:
            raise HTTPException(
                status_code=413,
                detail="文本文件超过 512KB，请拆分后再上传",
            )
        blocks = [{"type": "text", "text": f"【文件：{name}】\n{decoded}"}]
    elif suffix == ".docx":
        try:
            extracted = _docx_to_text(raw)
        except Exception:
            raise HTTPException(
                status_code=415,
                detail="docx 解析失败：文件可能损坏或不是有效的 Word 文档，"
                "请用 Word 重新另存为 .docx 后重试",
            ) from None
        if not extracted.strip():
            raise HTTPException(
                status_code=415,
                detail="docx 中没有可提取的文本（可能是纯图片文档）",
            )
        if len(extracted) > _TEXT_INLINE_LIMIT:
            raise HTTPException(
                status_code=413,
                detail="docx 提取出的文本超过 512KB，请拆分后重试",
            )
        blocks = [
            {
                "type": "text",
                "text": f"【文件：{name}（Word 文档，已提取正文文本）】\n{extracted}",
            }
        ]
    elif suffix == ".xlsx":
        try:
            extracted = _xlsx_to_text(raw)
        except Exception:
            raise HTTPException(
                status_code=415,
                detail="xlsx 解析失败：文件可能损坏或不是有效的 Excel "
                "文档，请用 Excel 重新另存为 .xlsx 后重试",
            ) from None
        if not extracted.strip():
            raise HTTPException(
                status_code=415,
                detail="xlsx 中没有可提取的数据（所有工作表均为空）",
            )
        if len(extracted) > _TEXT_INLINE_LIMIT:
            raise HTTPException(
                status_code=413,
                detail="xlsx 提取出的内容超过 512KB，请拆分工作表后再上传",
            )
        blocks = [
            {
                "type": "text",
                "text": f"【文件：{name}（Excel 表格，已提取为文本表格）】\n{extracted}",
            }
        ]
    elif (
        media.startswith("image/")
        or suffix in _IMAGE_EXTS
        or suffix == ".pdf"
    ):
        if suffix == ".pdf" and not media:
            media = "application/pdf"
        blocks = [
            {
                "type": "data",
                "source": {
                    "type": "base64",
                    "data": base64.b64encode(raw).decode("ascii"),
                    "media_type": media or "application/octet-stream",
                },
                "name": name,
            }
        ]
    else:
        raise HTTPException(
            status_code=415,
            detail="暂不支持该格式（模型层会丢弃该二进制内容）："
            "请转为 md/txt/csv/docx/xlsx（文本）、png/jpg（图片）或 pdf；"
            "旧版 .doc / .xls 请另存或导出为受支持格式",
        )

    message = (text or "").strip() or f"我上传了文件「{name}」，请查收。"
    session_id = await _relay_chat(db, run, project, message, blocks)
    return {"ok": True, "session_id": session_id, "file": name}


@router.post("/stage-runs/{run_id}/confirm")
async def confirm_stage_tools(
    run_id: int,
    payload: ConfirmIn,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> dict:
    """Answer the agent's human-in-the-loop tool permission prompt.

    The runtime parks a reply whenever a tool call needs a human
    verdict (session status ``awaiting_permission``); this resumes it.
    """
    run, project = await _get_run_or_404(db, run_id, user)
    if run.session_id is None:
        raise HTTPException(
            status_code=409, detail="该运行还没有智能体会话，无待确认工具"
        )
    client = get_runtime_client()
    try:
        agent_id = await flow.resolve_stage_agent(run.stage)
        # Pre-flight against the authoritative status: the persisted
        # message projection can lag (a just-confirmed call still reads
        # "asking"), and sending a stale confirm makes the runtime kill
        # the whole reply with a generic error. Only a parked session
        # accepts a confirm; anything else gets a clear 409.
        status = await client.get_session_status(agent_id, run.session_id)
        if status != "awaiting_permission":
            detail = (
                "智能体正在处理，请等它停靠后再确认"
                if status == "running"
                else "没有等待确认的工具调用"
            )
            raise HTTPException(status_code=409, detail=detail)
        data = await client.get_messages(agent_id, run.session_id)
    except (FlowError, RuntimeCallError) as exc:
        raise _flow_http(exc) from exc

    # The parked asking blocks may sit several messages back: failed
    # replies append empty error bubbles after them, and those carry
    # no tool calls. Scan backwards for the newest message that still
    # holds asking blocks instead of only looking at the last one.
    asking: list[dict] = []
    asking_msg: dict | None = None
    for msg in reversed(data["messages"]):
        asking = [
            b
            for b in (msg.get("content") or [])
            if isinstance(b, dict)
            and b.get("type") == "tool_call"
            and b.get("state") == "asking"
            and (
                not payload.tool_call_ids or b.get("id") in payload.tool_call_ids
            )
        ]
        if asking:
            asking_msg = msg
            break
    if asking_msg is None:
        raise HTTPException(
            status_code=409, detail="没有等待确认的工具调用"
        )
    # "全部允许"：停靠时把会话切到 accept_edits（运行中会被运行时
    # 409 拒绝——此时退回仅确认本次，下一轮停靠再试即可）。
    allow_all_applied = False
    if payload.allow_all and payload.confirmed:
        try:
            await client.set_session_permission_mode(
                agent_id, run.session_id, "accept_edits"
            )
            allow_all_applied = True
        except RuntimeCallError:
            pass
    try:
        await client.confirm_tool_calls(
            agent_id, run.session_id, asking_msg["id"], asking, payload.confirmed
        )
    except (FlowError, RuntimeCallError) as exc:
        raise _flow_http(exc) from exc
    return {
        "ok": True,
        "confirmed": payload.confirmed,
        "tool_calls": [b["name"] for b in asking],
        "allow_all_applied": allow_all_applied,
    }


@router.post("/stage-runs/{run_id}/interrupt")
async def interrupt_stage_run(
    run_id: int,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> dict:
    """Abort the run's in-flight or HITL-parked reply.

    The escape hatch for a parked session that can no longer be
    answered (e.g. its confirmation was lost to a failed reply): the
    parked tool calls close as interrupted and the session returns to
    idle, ready for the next message.
    """
    run, project = await _get_run_or_404(db, run_id, user)
    if run.session_id is None:
        raise HTTPException(
            status_code=409, detail="该运行还没有智能体会话，无可中止的回复"
        )
    client = get_runtime_client()
    try:
        agent_id = await flow.resolve_stage_agent(run.stage)
        await client.interrupt_session(agent_id, run.session_id)
    except (FlowError, RuntimeCallError) as exc:
        raise _flow_http(exc) from exc
    return {"ok": True, "session_id": run.session_id}


@router.get("/stage-runs/{run_id}/messages", response_model=MessagesOut)
async def stage_run_messages(
    run_id: int,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> MessagesOut:
    """Poll the run's conversation (frontend pull; SSE comes later)."""
    run, project = await _get_run_or_404(db, run_id, user)
    if run.session_id is None:
        return MessagesOut(messages=[], is_running=False)
    client = get_runtime_client()
    try:
        agent_id = await flow.resolve_stage_agent(run.stage)
        data = await client.get_messages(agent_id, run.session_id)
    except (FlowError, RuntimeCallError) as exc:
        raise _flow_http(exc) from exc
    return MessagesOut(**data)


@router.post("/stage-runs/{run_id}/submit-latest", response_model=ArtifactOut)
async def submit_latest(
    run_id: int,
    payload: SubmitLatestIn,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> ArtifactOut:
    """Submit the agent's latest answer as the stage artifact.

    The M1 stand-in for artifact extraction: title from the user,
    content = last non-empty assistant message of the session.
    """
    run, project = await _get_run_or_404(db, run_id, user)
    if run.session_id is None:
        raise HTTPException(
            status_code=409,
            detail="该运行还没有智能体会话（runtime 会话未创建），"
            "无可提交的产出",
        )
    client = get_runtime_client()
    try:
        agent_id = await flow.resolve_stage_agent(run.stage)
        data = await client.get_messages(agent_id, run.session_id)
    except (FlowError, RuntimeCallError) as exc:
        raise _flow_http(exc) from exc

    content = _last_assistant_text(data["messages"])
    if content is None:
        raise HTTPException(
            status_code=409,
            detail="会话中还没有智能体产出，先完成对话再提交",
        )
    try:
        artifact = await flow.submit_run(db, run, project, payload.title, content)
    except FlowError as exc:
        raise _flow_http(exc) from exc
    return ArtifactOut(
        id=artifact.id,
        version_no=artifact.version_no,
        title=artifact.title,
        format=artifact.format.value,
        storage_path=artifact.storage_path,
    )


def _last_assistant_text(messages: list[dict]) -> str | None:
    """Text of the last assistant message, tolerating both content
    shapes the runtime emits (plain string vs list of blocks)."""
    for msg in reversed(messages):
        if msg.get("role") != "assistant":
            continue
        raw = msg.get("content")
        if isinstance(raw, str):
            text = raw.strip()
        else:
            text = "\n".join(
                block.get("text", "")
                for block in raw or []
                if isinstance(block, dict) and block.get("type") == "text"
            ).strip()
        if text:
            return text
    return None


# ---------- workspace-file submit (dropdown picker) ----------

# What the picker offers; everything else in the workspace is scratch.
WORKSPACE_FILE_EXTS = {".md", ".html", ".txt", ".csv"}
_WORKSPACE_LIST_DEPTH = 2  # root + one subdirectory level
_WORKSPACE_LIST_CAP = 100


def _validate_workspace_path(path: str) -> None:
    if (
        ".." in path
        or "\\" in path
        or path.startswith("/")
        or Path(path).suffix.lower() not in WORKSPACE_FILE_EXTS
    ):
        raise HTTPException(
            status_code=400,
            detail=f"仅支持工作区内 {sorted(WORKSPACE_FILE_EXTS)} 文件，"
            f"路径需为相对路径且不含 ..",
        )


@router.get(
    "/stage-runs/{run_id}/workspace-files", response_model=list[WorkspaceFileOut]
)
async def list_workspace_files(
    run_id: int,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> list[WorkspaceFileOut]:
    """Candidate deliverable files in the run's session workspace.

    The runtime lists one directory level per call, so this walks root
    plus one subdirectory level (agents group their output in task
    folders) and filters by extension whitelist.
    """
    run, _project = await _get_run_or_404(db, run_id, user)
    if run.session_id is None:
        return []
    client = get_runtime_client()
    try:
        agent_id = await flow.resolve_stage_agent(run.stage)
        root = await client.list_workspace_dir(agent_id, run.session_id)
    except (FlowError, RuntimeCallError) as exc:
        raise _flow_http(exc) from exc

    files: list[WorkspaceFileOut] = []

    def _take(entry: dict, prefix: str = "") -> None:
        name = entry.get("name", "")
        if not entry.get("is_dir") and Path(name).suffix.lower() in WORKSPACE_FILE_EXTS:
            files.append(
                WorkspaceFileOut(
                    path=f"{prefix}{name}",
                    name=name,
                    size_bytes=entry.get("size_bytes"),
                )
            )

    for entry in root:
        if len(files) >= _WORKSPACE_LIST_CAP:
            break
        if entry.get("is_dir"):
            try:
                sub = await client.list_workspace_dir(
                    agent_id, run.session_id, entry.get("name", "")
                )
            except RuntimeCallError:
                continue  # unreadable folder — not worth failing the list
            for sub_entry in sub[: _WORKSPACE_LIST_CAP - len(files)]:
                _take(sub_entry, prefix=f"{entry.get('name', '')}/")
        else:
            _take(entry)
    return files


@router.get("/stage-runs/{run_id}/workspace-file")
async def get_workspace_file(
    run_id: int,
    path: str,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> dict:
    """Preview one workspace file BEFORE submitting it — the dropdown
    only shows names, this lets the user check the content first."""
    run, _project = await _get_run_or_404(db, run_id, user)
    if run.session_id is None:
        raise HTTPException(
            status_code=409, detail="该运行还没有智能体会话，无可预览的文件"
        )
    _validate_workspace_path(path)
    client = get_runtime_client()
    try:
        agent_id = await flow.resolve_stage_agent(run.stage)
        content = await client.read_workspace_file(
            agent_id, run.session_id, path
        )
    except (FlowError, RuntimeCallError) as exc:
        raise _flow_http(exc) from exc
    fmt = "html" if Path(path).suffix.lower() == ".html" else "md"
    return {"path": path, "format": fmt, "content": content}


@router.post("/stage-runs/{run_id}/submit-file", response_model=ArtifactOut)
async def submit_workspace_file(
    run_id: int,
    payload: SubmitFileIn,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> ArtifactOut:
    """Submit one workspace file (picked from the dropdown) as the
    stage artifact — the real deliverable the agent wrote, not the
    chat's last message."""
    run, project = await _get_run_or_404(db, run_id, user)
    if run.session_id is None:
        raise HTTPException(
            status_code=409,
            detail="该运行还没有智能体会话（runtime 会话未创建），"
            "无可提交的产出",
        )
    _validate_workspace_path(payload.path)
    client = get_runtime_client()
    try:
        agent_id = await flow.resolve_stage_agent(run.stage)
        content = await client.read_workspace_file(
            agent_id, run.session_id, payload.path
        )
    except (FlowError, RuntimeCallError) as exc:
        raise _flow_http(exc) from exc
    if not content or not content.strip():
        raise HTTPException(
            status_code=409, detail="选中的文件为空，无法作为产物提交"
        )
    suffix = Path(payload.path).suffix.lower()
    fmt = (
        ArtifactFormat.HTML if suffix == ".html" else ArtifactFormat.MD
    )
    try:
        artifact = await flow.submit_run(
            db, run, project, payload.title, content, fmt
        )
    except FlowError as exc:
        raise _flow_http(exc) from exc
    return ArtifactOut(
        id=artifact.id,
        version_no=artifact.version_no,
        title=artifact.title,
        format=artifact.format.value,
        storage_path=artifact.storage_path,
    )


# ---------- approvals ----------

@router.get("/approvals/pending", response_model=list[ApprovalOut])
async def pending_approvals(
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> list[ApprovalOut]:
    tasks = (
        await db.scalars(
            select(ApprovalTask)
            .options(selectinload(ApprovalTask.stage_run))
            .where(
                ApprovalTask.approver_user_id == user.id,
                ApprovalTask.status == ApprovalStatus.PENDING,
            )
            .order_by(ApprovalTask.id)
        )
    ).all()
    # 审批中心是跨项目的收件箱：把项目名一并带出，免得前端再逐个查
    project_names = dict(
        (
            await db.execute(
                select(Project.id, Project.name).where(
                    Project.id.in_({t.stage_run.project_id for t in tasks})
                )
            )
        ).all()
    )
    return [
        ApprovalOut(
            id=t.id,
            stage_run_id=t.stage_run_id,
            project_id=t.stage_run.project_id,
            project_name=project_names[t.stage_run.project_id],
            stage=t.stage_run.stage,
            run_round=t.stage_run.round,
            run_status=t.stage_run.status,
            status=t.status,
            comment=t.comment,
        )
        for t in tasks
    ]


@router.post("/approvals/{task_id}/decide")
async def decide_approval(
    task_id: int,
    payload: DecideIn,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> dict:
    task = await db.get(ApprovalTask, task_id)
    if task is None:
        raise HTTPException(status_code=404, detail="approval task not found")
    if task.approver_user_id != user.id:
        raise HTTPException(
            status_code=403, detail="only the assigned approver can decide"
        )
    run = await db.get(StageRun, task.stage_run_id)
    project = await db.get(Project, run.project_id)
    try:
        await flow.decide(db, task, run, project, payload.action, payload.comment)
    except FlowError as exc:
        raise _flow_http(exc) from exc
    return {
        "task_id": task.id,
        "action": payload.action,
        "run_status": run.status.value,
        "project_stage": project.current_stage.value,
        "project_status": project.status.value,
    }
