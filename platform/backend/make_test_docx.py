# -*- coding: utf-8 -*-
"""Create a realistic bank-requirement .docx for the upload e2e test."""
from docx import Document

doc = Document()
doc.add_heading("信贷系统需求书（测试）", level=1)
doc.add_heading("1 业务背景", level=2)
doc.add_paragraph("银行信贷审批流程目前为线下纸质流转，需要建设线上化系统。")
doc.add_heading("2 功能清单", level=2)
table = doc.add_table(rows=3, cols=3)
table.style = "Table Grid"
cells = [
    ["编号", "功能点", "优先级"],
    ["FR-01", "贷款申请录入", "高"],
    ["FR-02", "审批流程配置", "中"],
]
for i, row in enumerate(cells):
    for j, text in enumerate(row):
        table.rows[i].cells[j].text = text
doc.add_paragraph("以上功能需覆盖个人经营贷与小微信贷两类业务。")
doc.save("upload-test.docx")
print("created upload-test.docx")
